#!/usr/bin/env python3
"""Assemble Markdown sections into a Word document with native OMML equations.

Inputs
    --sections DIR     section sources, every *.md file in name order (or the order given by --order)
    --meta FILE        JSON with title_zh, title_en, optional subtitle, references (list of strings) and
                       glossary (list of {en, zh, def}); the format of an earlier writing workflow's result file
    --template FILE    optional .docx whose page setup, grid and styles are inherited; without it an A4
                       template with a line grid is created
    --figs FILE        optional JSON {"section file name": [png path, "图1.1 caption"], ...}
    --out FILE         output .docx
    --toc              add a table-of-contents field after the title

Markdown supported: # headings (levels 1 to 4), paragraphs, numbered and bulleted items, > quotes, pipe tables,
**bold**, `code`, inline $...$ and display $$...$$ math. Math goes through latex2mathml and mathml2omml and becomes
a native Word equation. A formula the converters cannot handle falls back to an image rendered by xelatex.

Known pitfalls, all handled below
    * mathml2omml closes <m:groupChrPr> with </m:groupChr> (\\underbrace, \\bar), fix_omml repairs it
    * a multiple line spacing on paragraphs snaps to the document grid and doubles the spacing, so none is set
    * first-line indent uses firstLineChars (2 characters), never spaces, and table cells get no indent
    * headings are left aligned and clamped to Heading 1..4, never the centred Title style
    * the reference list is numbered [1] [2] ...
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import struct
import subprocess
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor
from lxml import etree

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"


def fix_omml(omml: str) -> str:
    return re.sub(r"(<m:groupChrPr\b[^>]*>.*?)</m:groupChr>", r"\1</m:groupChrPr>", omml, flags=re.S)


def to_omml(latex: str) -> str:
    import latex2mathml.converter as l2m
    from mathml2omml import convert as m2o
    return fix_omml(m2o(l2m.convert(latex)))


def png_size(p):
    with open(p, "rb") as f:
        f.read(16)
        return struct.unpack(">II", f.read(8))


class Builder:
    def __init__(self, template: Path | None, cache: Path):
        self.cache = cache
        cache.mkdir(parents=True, exist_ok=True)
        self.doc = Document(str(template)) if template else self.blank_a4()
        body = self.doc.element.body
        self.sectPr = body.find(qn("w:sectPr"))
        for ch in list(body):
            if ch is not self.sectPr:
                body.remove(ch)
        self.styles()

    # ------------------------------------------------------------ template and styles
    @staticmethod
    def blank_a4():
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(25)
        sec.top_margin = sec.bottom_margin = Mm(25)
        grid = sec._sectPr.find(qn("w:docGrid"))
        if grid is None:
            grid = OxmlElement("w:docGrid")
            sec._sectPr.append(grid)
        grid.set(qn("w:type"), "lines")
        grid.set(qn("w:linePitch"), "312")
        return doc

    @staticmethod
    def set_ea(style, west, ea, size, bold=None, black=False):
        style.font.name = west
        style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if black:
            style.font.color.rgb = RGBColor(0, 0, 0)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), ea)
        rf.set(qn("w:ascii"), west)
        rf.set(qn("w:hAnsi"), west)

    def styles(self):
        st = self.doc.styles
        self.set_ea(st["Normal"], "Times New Roman", "宋体", 12)
        st["Normal"].paragraph_format.space_after = Pt(2)
        self.set_ea(st["Heading 1"], "宋体", "宋体", 16, bold=True, black=True)
        for h in ("Heading 2", "Heading 3", "Heading 4"):
            self.set_ea(st[h], "Times New Roman", "宋体", 12, bold=True, black=True)
        for h in ("Heading 1", "Heading 2", "Heading 3", "Heading 4"):
            pf = st[h].paragraph_format
            pf.space_before, pf.space_after = Pt(8), Pt(4)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # ------------------------------------------------------------ paragraph helpers
    @staticmethod
    def indent(p, chars=200, twips="480"):
        pPr = p._p.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn("w:firstLineChars"), str(chars))
        ind.set(qn("w:firstLine"), twips)

    def noindent(self, p):
        self.indent(p, 0, "0")

    @staticmethod
    def run_ea(r, west="宋体", ea="宋体"):
        rf = r._element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:eastAsia"), ea)
        rf.set(qn("w:ascii"), west)
        rf.set(qn("w:hAnsi"), west)

    def heading(self, text, lvl):
        lvl = max(1, min(lvl, 4))
        p = self.doc.add_heading(text, level=lvl)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            self.run_ea(r, "宋体" if lvl == 1 else "Times New Roman", "宋体")
        return p

    def formula_image(self, latex, display):
        key = hashlib.md5((("D" if display else "I") + latex).encode()).hexdigest()[:16]
        png = self.cache / f"{key}.png"
        if png.exists():
            return png
        inner = ("\\[" + latex + "\\]") if display else ("$" + latex + "$")
        (self.cache / f"{key}.tex").write_text(
            "\\documentclass[border=1pt]{standalone}\n\\usepackage{ctex}\\usepackage{amsmath}\\usepackage{amssymb}\n"
            "\\begin{document}" + inner + "\\end{document}\n", encoding="utf-8")
        subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", f"{key}.tex"], cwd=self.cache,
                       capture_output=True)
        subprocess.run(["pdftoppm", "-png", "-r", "340", "-singlefile", f"{key}.pdf", key], cwd=self.cache,
                       capture_output=True)
        return png if png.exists() else None

    def runs(self, p, text, size=None):
        for tk in re.split(r"(\$[^$\n]+?\$|\*\*.+?\*\*|`[^`]+?`)", text):
            if not tk:
                continue
            if len(tk) > 2 and tk[0] == "$" and tk[-1] == "$":
                latex = tk[1:-1]
                try:
                    p._p.append(etree.fromstring('<r xmlns:m="%s">%s</r>' % (M_NS, to_omml(latex)))[0])
                except Exception:
                    img = self.formula_image(latex, False)
                    if img:
                        p.add_run().add_picture(str(img), height=Pt(13))
                    else:
                        r = p.add_run(latex)
                        if size:
                            r.font.size = Pt(size)
                continue
            if tk.startswith("**") and tk.endswith("**"):
                r = p.add_run(tk[2:-2])
                r.bold = True
            elif tk.startswith("`") and tk.endswith("`"):
                r = p.add_run(tk[1:-1])
                r.font.name = "Consolas"
            else:
                r = p.add_run(tk)
            if size:
                r.font.size = Pt(size)

    def display(self, latex):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            p._p.append(etree.fromstring('<m:oMathPara xmlns:m="%s">%s</m:oMathPara>' % (M_NS, to_omml(latex))))
        except Exception:
            img = self.formula_image(latex, True)
            if img:
                p.add_run().add_picture(str(img), width=Inches(min(5.5, png_size(img)[0] / 340.0)))
            else:
                p.add_run(latex)

    def table(self, header, rows):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for k, c in enumerate(header):
            cell = t.rows[0].cells[k]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self.noindent(cell.paragraphs[0])
            self.runs(cell.paragraphs[0], c, size=10.5)
            for run in cell.paragraphs[0].runs:
                run.bold = True
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "E7E7E7")
            cell._tc.get_or_add_tcPr().append(shd)
        for row in rows:
            cs = t.add_row().cells
            for k in range(len(header)):
                self.noindent(cs[k].paragraphs[0])
                self.runs(cs[k].paragraphs[0], row[k] if k < len(row) else "", size=10.5)

    def figure(self, png, caption, maxw=6.5):
        w, h = png_size(png)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if w >= h:
            p.add_run().add_picture(str(png), width=Inches(min(maxw, w / 200.0)))
        else:
            p.add_run().add_picture(str(png), height=Inches(4.2))
        c = self.doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        m = re.match(r"(图[0-9.]+|Figure\s*[0-9.]+)(.*)", caption)
        if m:
            r = c.add_run(m.group(1) + " ")
            r.bold = True
            r.font.size = Pt(10.5)
            c.add_run(m.group(2).strip()).font.size = Pt(10.5)
        else:
            c.add_run(caption).font.size = Pt(10.5)

    def toc(self):
        p = self.doc.add_paragraph()
        r = p.add_run()
        for kind, text in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None),
                           (None, "在 Word 中右键此处选择更新域以生成目录"), ("end", None)):
            if kind:
                e = OxmlElement("w:fldChar")
                e.set(qn("w:fldCharType"), kind)
                r._r.append(e)
            elif text.startswith("TOC"):
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r._r.append(it)
            else:
                t = OxmlElement("w:t")
                t.text = text
                r._r.append(t)

    # ------------------------------------------------------------ markdown
    def render(self, md):
        md = md.replace("&gt;", ">").replace("&lt;", "<")
        lines = md.split("\n")
        i, n = 0, len(lines)
        is_sep = lambda ln: bool(re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", ln)) and "-" in ln  # noqa: E731
        while i < n:
            ln = lines[i]
            s = ln.strip()
            if s.startswith("$$"):
                buf = ln
                while buf.count("$$") < 2 and i + 1 < n:
                    i += 1
                    buf += "\n" + lines[i]
                m = re.search(r"\$\$(.+?)\$\$", buf, re.S)
                if m:
                    self.display(m.group(1).strip())
                i += 1
                continue
            if s.startswith("|") and i + 1 < n and is_sep(lines[i + 1]):
                header = [c.strip() for c in s.strip("|").split("|")]
                rows, j = [], i + 2
                while j < n and lines[j].strip().startswith("|"):
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                    j += 1
                self.table(header, rows)
                i = j
                continue
            m = re.match(r"^(#{1,6})\s+(.*)$", ln)
            if m:
                self.heading(m.group(2).strip(), len(m.group(1)))
            elif re.match(r"^\s*(\d+)[.)]\s+", ln):
                m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", ln)
                p = self.doc.add_paragraph()
                self.indent(p)
                self.runs(p, f"{m.group(1)}. {m.group(2)}")
            elif re.match(r"^\s*[-*]\s+", ln):
                p = self.doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(12)
                self.runs(p, "· " + re.sub(r"^\s*[-*]\s+", "", ln))
            elif s.startswith(">"):
                p = self.doc.add_paragraph()
                self.runs(p, s[1:].strip())
                for r in p.runs:
                    r.italic = True
                    r.font.size = Pt(10.5)
            elif s and not re.match(r"^\s*---+\s*$", ln):
                p = self.doc.add_paragraph()
                self.indent(p)
                self.runs(p, s)
            i += 1

    def save(self, out: Path):
        body = self.doc.element.body
        body.remove(self.sectPr)
        body.append(self.sectPr)
        self.doc.save(str(out))


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--sections", required=True, type=Path)
    ap.add_argument("--meta", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--template", type=Path)
    ap.add_argument("--figs", type=Path)
    ap.add_argument("--order", default="", help="comma separated section file names, default all *.md in name order")
    ap.add_argument("--toc", action="store_true")
    a = ap.parse_args()

    meta = json.loads(a.meta.read_text(encoding="utf-8"))
    figs = json.loads(a.figs.read_text(encoding="utf-8")) if a.figs else {}
    files = [a.sections / f for f in a.order.split(",") if f] or sorted(a.sections.glob("*.md"))
    b = Builder(a.template, Path(tempfile.gettempdir()) / "build_docx_formula_cache")

    p = b.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("title_zh", ""))
    r.bold, r.font.size = True, Pt(16)
    b.run_ea(r)
    if meta.get("title_en"):
        p = b.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta["title_en"])
        r.bold, r.font.size, r.font.name = True, Pt(13), "Times New Roman"
    if meta.get("subtitle"):
        p = b.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(meta["subtitle"])
        r.italic, r.font.size = True, Pt(10.5)
    if a.toc:
        b.heading("目录", 1)
        b.toc()

    for f in files:
        b.render(f.read_text(encoding="utf-8"))
        if f.name in figs:
            png, cap = figs[f.name]
            b.figure(Path(png) if os.path.isabs(png) else a.figs.parent / png, cap)
    refs = meta.get("references", [])
    if refs:
        b.heading("参考文献", 1)
        for i, ref in enumerate(refs, 1):
            p = b.doc.add_paragraph()
            b.noindent(p)
            b.runs(p, f"[{i}] {ref}")
    glossary = meta.get("glossary", [])
    if glossary:
        b.heading("附录 术语表", 1)
        b.table(["英文", "中文", "说明"], [[g.get("en", ""), g.get("zh", ""), g.get("def", "")] for g in glossary])
    b.save(a.out)
    xml = zipfile.ZipFile(a.out).read("word/document.xml").decode("utf-8", "ignore")
    print("saved", a.out)
    print("  tables:", xml.count("<w:tbl>"), "| images:", xml.count("<a:blip"), "| native equations:",
          xml.count("<m:oMath"), "| first-line indented paragraphs:", xml.count('w:firstLineChars="200"'),
          "| TOC field:", "TOC" in xml)


if __name__ == "__main__":
    main()
