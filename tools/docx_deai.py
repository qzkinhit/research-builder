#!/usr/bin/env python3
"""docx 去 AI 怪词。把 Word 稿里的生硬词、口语比喻词、元话术和禁用标点就地换掉，保留格式。
run 级替换保留每段的粗体与字体。短语规则在前，先处理「主张」的动词、序号等上下文。
规则见 writing-deai.md 第 2 节与第 29 条。替换只能清词，句式与结构还要按 writing-deai.md 第 3 节人工改。

  python docx_deai.py <file.docx>            # dry-run: 只报会改哪些(不落盘)
  python docx_deai.py <file.docx> --apply    # 落盘(先备份 .bak), 再自检建议跑 polish_check

需要 python-docx。
"""
import sys, shutil

# 有序替换(长/带上下文的在前)。中文短语优先, 再到单词。
REPL = [
    # 主张: 动词/序号/名词分开("主张检测"在前, 让"不再主张检测"->"不再以检测为目标")
    ("主张检测", "以检测为目标"), ("本文主张的是", "本文的论点是"), ("我们主张", "我们认为"),
    ("主张三", "论点三"), ("主张二", "论点二"), ("主张一", "论点一"),
    ("核心主张", "核心论点"), ("主张是", "论点是"), ("主张的", "论点的"),
    ("主张", "论点"),
    # 口径(常见搭配在前)
    ("统计口径", "统计方式"), ("评测口径", "评测方式"), ("口径统一", "表述统一"),
    ("口径下调", "改用更保守的表述"), ("统一口径", "统一表述"), ("口径", "表述"),
    # AI 生硬词
    ("诚实边界", "局限性"), ("防质疑审计", "防质疑核查"), ("审计", "核查"),
    ("平局组", "并列梯队"), ("平局", "并列"), ("证据组", "证据"),
    ("稳健性", "鲁棒性"), ("稳健", "鲁棒"),
    # 口语/比喻词
    ("通吃", "通用"), ("崩盘", "大幅退化"), ("冷宫", "最差"), ("王座", "最优"),
    ("是王", "最优"), ("垫底", "最差"), ("死穴", "症结"), ("翻盘", "反超"),
    ("卖点", "核心贡献"), ("无敌", "严谨"), ("打脸", "被推翻"),
    ("抢光", "挤占"), ("饿死", "得不到资源"),
    # 更花哨的名次/涨跌比喻(长的在前)
    ("落入末位", "居末"), ("末位区间", "最差区间"), ("末位", "最差"),
    ("夺冠", "最优"), ("称王", "最优"), ("登顶", "居首"), ("问鼎", "居首"), ("称霸", "最优"),
    ("碾压", "大幅领先"), ("吊打", "大幅领先"), ("完胜", "显著优于"), ("秒杀", "大幅领先"),
    ("跌破", "低于"), ("跌入", "落到"), ("反超", "超过"), ("驾驶位", "主导"),
    ("翻车", "失败"), ("拉胯", "表现差"), ("用武之地", "适用场景"), ("有的放矢", "有针对性"), ("遥遥领先", "显著领先"), ("一骑绝尘", "显著领先"),
    # AI 口癖/元话术: 删掉"先宣布再说"那层(连尾逗号一起删)
    ("换言之，", ""), ("换言之,", ""), ("换言之", "即"),
    ("综上所述，", ""), ("综上所述", ""), ("总而言之，", ""), ("总的来说，", ""), ("总体而言，", ""),
    ("值得注意的是，", ""), ("值得注意的是,", ""), ("值得注意的是 ", ""), ("值得注意的是", ""),
    ("值得一提的是，", ""), ("值得强调的是，", ""),
    ("需要指出的是，", ""), ("需要强调的是，", ""), ("需要说明的是，", ""), ("需要注意的是，", ""),
    ("需要诚实指出，", ""), ("需要诚实指出", ""),
    ("我们诚实地承认", "我们承认"), ("诚实地承认", "承认"), ("诚实地指出", "指出"),
    ("我们诚实地", "我们"), ("不难看出，", ""), ("不难发现，", ""),
    # 分号 -> 逗号(铁律: 用句号逗号, 不用分号)
    ("；", "，"), (";", ","),
    # 花引号 -> 直引号(铁律: 不用花引号)
    ("\u201c", '"'), ("\u201d", '"'), ("\u2018", "'"), ("\u2019", "'"),
]


def do_run(text, counts=None):
    for a, b in REPL:
        if a in text:
            if counts is not None:
                counts[a] = counts.get(a, 0) + text.count(a)
            text = text.replace(a, b)
    return text


def main():
    if len(sys.argv) < 2:
        print("用法: python docx_deai.py <file.docx> [--apply]"); return 2
    import docx
    path = sys.argv[1]; apply = "--apply" in sys.argv
    d = docx.Document(path)
    counts = {}

    def handle_paras(paras):
        for p in paras:
            runs = p.runs
            if not runs:
                continue
            full = "".join(r.text for r in runs)
            target = do_run(full, counts)       # 计数只在段落级做一次
            if not apply or target == full:
                continue
            for r in runs:                      # 先 run 级替换, 保留每段的粗体与字体
                r.text = do_run(r.text)
            if "".join(r.text for r in runs) != target:
                runs[0].text = target           # 兜底: 被拆在多个 run 里的词或短语
                for r in runs[1:]:
                    r.text = ""

    handle_paras(d.paragraphs)
    for tb in d.tables:
        for row in tb.rows:
            for cell in row.cells:
                handle_paras(cell.paragraphs)

    if not counts:
        print("无怪词, 无需改"); return 0
    total = sum(counts.values())
    print(f"{'[已改]' if apply else '[dry-run 会改]'} 共 {total} 处:")
    for a in sorted(counts, key=lambda k: -counts[k]):
        b = next(v for k, v in REPL if k == a)
        print(f"  {a} -> {b} : {counts[a]} 处")
    if apply:
        shutil.copy(path, path + ".bak")
        d.save(path)
        print(f"已落盘(备份 {path}.bak). 建议跑: python tools/polish_check.py {path}")
    else:
        print("dry-run, 未落盘. 确认后加 --apply")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
